# modules/connectors/parsers/docx_parser.py
import io
from modules.connectors.parsers.base import BaseParser

class DocxParser(BaseParser):
    def parse(self, content_bytes: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n".join(full_text)
        except Exception:
            return content_bytes.decode("utf-8", errors="ignore")
