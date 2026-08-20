# app/parsers/docx_parser.py
from io import BytesIO
from app.parsers.base import BaseParser

class DocxParser(BaseParser):
    """
    Parser for Microsoft Word (.docx) documents using python-docx.
    """
    def parse(self, content_bytes: bytes) -> str:
        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is not installed. Run `pip install python-docx` to parse DOCX files.")

        doc = docx.Document(BytesIO(content_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    full_text.append(row_text)
        return "\n".join(full_text)
